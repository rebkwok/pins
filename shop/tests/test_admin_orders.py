"""
The orders admin (shop/admin_orders.py): the orders list, its filters and
export, the order page and its actions. Refunds are tested in test_refunds.py.
"""

import csv
import html
import io
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from django.contrib.auth.models import Permission
from django.contrib.messages import get_messages
from django.core import mail
from django.urls import reverse
from model_bakery import baker
from openpyxl import load_workbook
from salesman.core.utils import get_salesman_model

from ..admin_orders import OrderPermissionPolicy
from ..models import OrderItem
from ..status import OrderStatus
from ..templatetags.shoptags import gbp
from .test_shipping import dispatcher, make_dispatcher, make_order, shop_variants  # noqa: F401

pytestmark = pytest.mark.django_db

Order = get_salesman_model("Order")
OrderNote = get_salesman_model("OrderNote")

INDEX_URL = reverse("shop_orders:index")


@pytest.fixture
def admin_client(client):
    client.force_login(baker.make("auth.User", is_superuser=True, is_staff=True))
    yield client


def order_url(order, action=None):
    name = f"shop_orders:{action or 'edit'}"
    return reverse(name, args=(order.pk,))


def messages_for(resp):
    return [html.unescape(str(m)).strip() for m in get_messages(resp.wsgi_request)]


def paid_order(variants, **kwargs):
    order = make_order(variants, **kwargs)
    order.pay(amount=order.total, transaction_id=f"pi_{order.pk}", payment_method="stripe")
    mail.outbox.clear()
    return order


def listed(resp):
    return set(resp.context["object_list"])


# Plain statuses and amounts


def test_gbp():
    assert gbp(Decimal("5")) == "£5.00"
    assert gbp("-4.99") == "-£4.99"
    assert gbp(1234.5) == "£1,234.50"
    assert gbp(None) == ""


@pytest.mark.parametrize(
    "status,shipping_method,paid,expected",
    [
        (OrderStatus.NEW, "deliver", False, "Awaiting payment"),
        (OrderStatus.HOLD, "deliver", False, "Awaiting payment"),
        (OrderStatus.FAILED, "deliver", False, "Payment failed"),
        (OrderStatus.PROCESSING, "deliver", False, "Processing"),
        (OrderStatus.PROCESSING, "deliver", True, "Paid – ready to ship"),
        (OrderStatus.PROCESSING, "collect", True, "Paid – ready to collect"),
        (OrderStatus.PARTIALLY_SHIPPED, "deliver", True, "Partly shipped"),
        (OrderStatus.PARTIALLY_SHIPPED, "collect", True, "Partly collected"),
        (OrderStatus.SHIPPED, "deliver", True, "Shipped"),
        (OrderStatus.SHIPPED, "collect", True, "Collected"),
        (OrderStatus.COMPLETED, "deliver", True, "Completed"),
        (OrderStatus.CANCELLED, "deliver", False, "Cancelled"),
        (OrderStatus.REFUNDED, "deliver", True, "Refunded"),
    ],
)
def test_display_status(order, status, shipping_method, paid, expected):
    order.status = status
    order.shipping_method = shipping_method
    if paid:
        order.pay(amount=order.total, transaction_id="pi_1")
    assert order.display_status == expected
    # the customer's is the same, but a collection may not be ready yet
    if expected == "Paid – ready to collect":
        expected = "Paid – to be collected"
    assert order.customer_display_status == expected


def test_display_status_partly_refunded(order):
    order.status = OrderStatus.SHIPPED
    order.pay(amount=order.total, transaction_id="pi_1")
    order.pay(amount=Decimal("-5"), transaction_id="re_1")
    # (collect in person)
    assert order.display_status == "Collected (partly refunded)"


# Orders list


def test_orders_list(admin_client, shop_variants):
    order = paid_order(shop_variants)
    resp = admin_client.get(INDEX_URL)
    assert resp.status_code == 200
    content = resp.content.decode()
    assert order.ref in content
    assert order_url(order) in content
    assert f"£{order.total:.2f}" in content
    assert "Paid – ready to ship" in content
    assert "c@test.com" in content
    # orders can't be added
    assert reverse("shop_orders:index") + "new/" not in content


def test_orders_list_search(admin_client, shop_variants):
    order = make_order(shop_variants)
    order.name = "Maria Lopez"
    order.save()
    other = make_order(shop_variants)
    for query in ["maria", order.ref]:
        resp = admin_client.get(INDEX_URL, {"q": query})
        assert listed(resp) == {order}
    assert listed(admin_client.get(INDEX_URL, {"q": "c@test.com"})) == {order, other}


def test_orders_list_search_items(admin_client, shop_variants):
    # e.g. orders that included hoodies
    mug, card, raffle = shop_variants
    with_mug = make_order([mug, card])
    make_order([card])
    resp = admin_client.get(INDEX_URL, {"q": "mug"})
    # (listed once)
    assert list(resp.context["object_list"]) == [with_mug]
    two_mugs = make_order([mug])
    baker.make(
        "shop.OrderItem",
        order=two_mugs,
        product_data={"name": "Another mug"},
        quantity=1,
        unit_price=1,
        subtotal=1,
        total=1,
    )
    resp = admin_client.get(INDEX_URL, {"q": "mug"})
    assert list(resp.context["object_list"]) == [two_mugs, with_mug]


def test_orders_list_newest_first(admin_client, shop_variants):
    orders = [make_order(shop_variants) for _ in range(3)]
    resp = admin_client.get(INDEX_URL)
    assert list(resp.context["object_list"]) == orders[::-1]


@pytest.fixture
def orders_by_status(shop_variants):
    return {
        status: make_order(shop_variants, status=status)
        for status in [
            OrderStatus.HOLD,
            OrderStatus.FAILED,
            OrderStatus.PROCESSING,
            OrderStatus.PARTIALLY_SHIPPED,
            OrderStatus.SHIPPED,
            OrderStatus.COMPLETED,
            OrderStatus.CANCELLED,
            OrderStatus.REFUNDED,
        ]
    }


@pytest.mark.parametrize(
    "value,statuses",
    [
        ("awaiting_payment", [OrderStatus.HOLD]),
        ("failed", [OrderStatus.FAILED]),
        ("processing", [OrderStatus.PROCESSING]),
        ("partly_shipped", [OrderStatus.PARTIALLY_SHIPPED]),
        ("shipped", [OrderStatus.SHIPPED]),
        ("completed", [OrderStatus.COMPLETED]),
        ("cancelled", [OrderStatus.CANCELLED]),
        ("refunded", [OrderStatus.REFUNDED]),
    ],
)
def test_orders_status_filter(admin_client, orders_by_status, value, statuses):
    resp = admin_client.get(INDEX_URL, {"status": value})
    assert listed(resp) == {orders_by_status[status] for status in statuses}


def test_orders_status_filter_choices(admin_client, orders_by_status):
    content = admin_client.get(INDEX_URL).content.decode()
    # salesman's unused statuses aren't offered
    assert ">Awaiting payment</option>" in content
    assert ">Created</option>" not in content
    assert ">New</option>" not in content


def test_orders_to_ship_filter(admin_client, shop_variants, dispatcher):
    mug, card, raffle = shop_variants
    to_ship = make_order([mug, raffle])
    all_shipped = make_order([mug])
    all_shipped.mark_items_shipped(all_shipped.items.all(), dispatcher)
    nothing_to_ship = make_order([raffle])  # completed
    unpaid = make_order([mug], status=OrderStatus.HOLD)
    assert nothing_to_ship.status == OrderStatus.COMPLETED
    assert listed(admin_client.get(INDEX_URL, {"to_ship": "true"})) == {to_ship}
    assert listed(admin_client.get(INDEX_URL, {"to_ship": "false"})) == {
        all_shipped,
        nothing_to_ship,
        unpaid,
    }


def test_orders_dispatcher_filter(admin_client, shop_variants, dispatcher):
    mug, card, raffle = shop_variants
    anns = make_order([mug, card])
    no_dispatcher = make_order([card])
    # a former dispatcher can still be chosen, for their past orders
    bob = make_dispatcher("Bob", "bob@test.com")
    bobs = make_order([card])
    OrderItem.objects.filter(order=bobs).update(dispatcher=bob, dispatcher_name="Bob Smith")
    bob.groups.clear()
    content = admin_client.get(INDEX_URL).content.decode()
    assert ">Ann Smith</option>" in content
    assert ">Bob Smith</option>" in content
    assert listed(admin_client.get(INDEX_URL, {"dispatcher": dispatcher.pk})) == {anns}
    assert listed(admin_client.get(INDEX_URL, {"dispatcher": bob.pk})) == {bobs}
    assert no_dispatcher not in listed(admin_client.get(INDEX_URL, {"dispatcher": dispatcher.pk}))


def test_orders_product_filter(admin_client, shop_variants):
    mug, card, raffle = shop_variants
    mugs = make_order([mug, raffle])
    cards = make_order([card])
    assert listed(admin_client.get(INDEX_URL, {"product": mug.product.pk})) == {mugs}
    assert listed(admin_client.get(INDEX_URL, {"product": card.product.pk})) == {cards}


@pytest.fixture
def refunded_orders(basket):
    # unpaid, paid, partially refunded and fully refunded orders (totals £20)
    orders = [Order.objects.create_from_basket(basket, request=None) for _ in range(4)]
    for order in orders[1:]:
        order.pay(amount=Decimal("20.00"), transaction_id=f"pi_{order.id}")
    orders[2].pay(amount=Decimal("-5.00"), transaction_id="re_1")
    orders[3].pay(amount=Decimal("-20.00"), transaction_id="re_2")
    yield orders


@pytest.mark.parametrize(
    "value,expected", [("partial", [2]), ("full", [3]), ("none", [0, 1]), ("", [0, 1, 2, 3])]
)
def test_orders_refunds_filter(admin_client, refunded_orders, value, expected):
    resp = admin_client.get(INDEX_URL, {"refunds": value})
    assert listed(resp) == {refunded_orders[index] for index in expected}


def test_orders_date_filter(admin_client, shop_variants):
    old, new = make_order(shop_variants), make_order(shop_variants)
    Order.objects.filter(pk=old.pk).update(date_created=datetime(2026, 1, 10, 12, tzinfo=UTC))
    Order.objects.filter(pk=new.pk).update(date_created=datetime(2026, 3, 10, 12, tzinfo=UTC))
    resp = admin_client.get(
        INDEX_URL, {"date_created_from": "2026-03-01", "date_created_to": "2026-03-31"}
    )
    assert listed(resp) == {new}


# Export: one row per order item


def csv_rows(resp):
    content = b"".join(resp.streaming_content).decode("utf-8-sig")
    return list(csv.DictReader(io.StringIO(content)))


def test_orders_export_csv(admin_client, shop_variants, dispatcher):
    mug, card, raffle = shop_variants
    order = paid_order([mug, card, raffle])
    order.name = "Cat"
    order.shipping_address = "1 Road\nTown"
    order.save()
    order.mark_items_shipped([order.items.get(product_id=mug.id)], dispatcher)
    make_order([card], status=OrderStatus.HOLD)  # filtered out
    resp = admin_client.get(INDEX_URL, {"status": "partly_shipped", "export": "csv"})
    assert resp.status_code == 200
    rows = csv_rows(resp)
    assert [row["Item"] for row in rows] == ["Test Product - Mug", "Card", "Raffle ticket"]
    mug_row = rows[0]
    assert mug_row["Order"] == order.ref
    assert mug_row["Customer"] == "Cat"
    assert mug_row["Email"] == "c@test.com"
    assert mug_row["Shipping method"] == "Delivery"
    assert mug_row["Shipping address"] == "1 Road\nTown"
    assert mug_row["Order status"] == "Partly shipped"
    assert mug_row["Quantity"] == "2"
    assert mug_row["Item total (£)"] == "20.00"
    assert mug_row["Dispatcher"] == "Ann Smith"
    assert mug_row["Item status"] == "Shipped/collected"
    # dates without seconds (the site's time zone is UTC)
    assert mug_row["Date"] == f"{order.date_created:%Y-%m-%d %H:%M}"
    assert len(mug_row["Shipped/collected"]) == len("2026-10-24 14:05")
    assert rows[1]["Shipped/collected"] == ""
    assert rows[1]["Dispatcher"] == ""
    assert rows[1]["Item status"] == "Ready to send"
    assert rows[2]["Item status"] == "No shipping needed"


def test_orders_export_dispatcher_items_only(admin_client, shop_variants, dispatcher):
    # filtered by dispatcher (or product), only their items are exported
    mug, card, raffle = shop_variants
    paid_order([mug, card])
    resp = admin_client.get(INDEX_URL, {"dispatcher": dispatcher.pk, "export": "csv"})
    assert [row["Item"] for row in csv_rows(resp)] == ["Test Product - Mug"]
    resp = admin_client.get(INDEX_URL, {"product": card.product.pk, "export": "csv"})
    assert [row["Item"] for row in csv_rows(resp)] == ["Card"]


def test_orders_export_ticked(admin_client, shop_variants):
    # only the ticked orders (still within any filters)
    mug, card, raffle = shop_variants
    first, second, third = paid_order([mug]), paid_order([card]), paid_order([mug])
    params = {"export": "csv", "order": [first.pk, third.pk, "x"]}
    rows = csv_rows(admin_client.get(INDEX_URL, params))
    assert [row["Order"] for row in rows] == [third.ref, first.ref]
    params["product"] = card.product.pk
    assert csv_rows(admin_client.get(INDEX_URL, params)) == []


def test_orders_export_filename(admin_client, shop_variants, freezer):
    freezer.move_to(datetime(2026, 10, 24, 14, 5, tzinfo=UTC))
    paid_order(shop_variants)
    resp = admin_client.get(INDEX_URL, {"export": "xlsx"})
    assert 'filename="orders-2026-10-24-1405.xlsx"' in resp["Content-Disposition"]


def test_orders_list_selection(admin_client, shop_variants):
    order = paid_order(shop_variants)
    content = admin_client.get(INDEX_URL).content.decode()
    assert f'name="order" value="{order.pk}"' in content
    assert f'aria-label="Select order {order.ref}, to download"' in content
    assert "data-order-select-all" in content
    assert "shop/js/order_export.js" in content
    # download buttons (disabled until orders are ticked) in the header, beside
    # the title, rather than in its menu
    header = content.split('class="w-slim-header')[1].split("</header>")[0]
    assert 'data-order-export="xlsx" disabled' in header
    assert 'data-order-export="csv" disabled' in content
    assert "Download (Excel)" in content
    assert "?export=xlsx" not in content
    assert "to download all 1." in " ".join(content.split())


def test_orders_list_selection_paginated(admin_client, shop_variants, settings):
    for _ in range(51):
        make_order(shop_variants)
    content = admin_client.get(INDEX_URL).content.decode()
    assert "to download all 51 (on every page)." in " ".join(content.split())


def test_orders_export_xlsx_orders_shaded(admin_client, shop_variants):
    # each order's rows are shaded in turn (grey, none, grey...)
    mug, card, raffle = shop_variants
    for _ in range(3):
        paid_order([mug, card])
    resp = admin_client.get(INDEX_URL, {"export": "xlsx"})
    sheet = load_workbook(io.BytesIO(b"".join(resp.streaming_content))).active
    fills = [
        row[0].fill.fgColor.rgb if row[0].fill.fill_type else None
        for row in sheet.iter_rows(min_row=2)
    ]
    grey = "00D9D9D9"
    assert fills == [grey, grey, None, None, grey, grey]
    # every cell in the row, not just the first
    assert {cell.fill.fill_type for cell in next(sheet.iter_rows(min_row=2))} == {"solid"}
    # bold headings (not shaded), and borders on every cell
    assert sheet["A1"].fill.fill_type is None
    assert all(cell.font.bold for cell in sheet[1])
    assert not sheet["A2"].font.bold
    assert {
        (cell.border.left.style, cell.border.bottom.style)
        for row in sheet.iter_rows()
        for cell in row
    } == {("thin", "thin")}


def test_orders_export_xlsx(admin_client, shop_variants):
    order = paid_order(shop_variants)
    resp = admin_client.get(INDEX_URL, {"export": "xlsx"})
    assert resp.status_code == 200
    sheet = load_workbook(io.BytesIO(b"".join(resp.streaming_content))).active
    rows = list(sheet.values)
    assert rows[0][:3] == ("Order", "Date", "Customer")
    assert len(rows) == 4
    assert rows[1][0] == order.ref


# Permissions


def test_orders_need_change_permission(client, order):
    user = baker.make("auth.User", is_staff=True)
    user.user_permissions.add(Permission.objects.get(codename="access_admin"))
    client.force_login(user)
    assert client.get(INDEX_URL).status_code == 302
    assert client.get(order_url(order)).status_code == 302
    user.user_permissions.add(Permission.objects.get(codename="change_order"))
    assert client.get(INDEX_URL).status_code == 200
    assert client.get(order_url(order)).status_code == 200


def test_order_permission_policy(order):
    user = baker.make("auth.User", is_superuser=True)
    policy = OrderPermissionPolicy(Order)
    assert policy.user_has_permission(user, "change")
    assert not policy.user_has_permission(user, "add")
    assert not policy.user_has_permission(user, "delete")
    assert policy.user_has_any_permission(user, ["add", "change"])
    assert not policy.user_has_any_permission(user, ["add", "delete"])


# The order page


def test_order_page(admin_client, shop_variants, dispatcher):
    mug, card, raffle = shop_variants
    order = paid_order([mug, card, raffle])
    order.name = "Cat"
    order.shipping_address = "1 Road\nTown"
    order.save()
    OrderNote.objects.create(order=order, message="Gift wrap please")
    resp = admin_client.get(order_url(order))
    assert resp.status_code == 200
    content = html.unescape(resp.content.decode())
    assert f"Order {order.ref}" in content
    assert "Paid – ready to ship" in content
    assert "Cat" in content
    assert 'href="mailto:c@test.com"' in content
    assert "1 Road<br>Town" in content
    # items, with their dispatchers and status
    assert "Test Product - Mug" in content
    assert "Ann Smith" in content
    assert "No shipping needed" in content
    assert content.count('name="items"') == 2
    assert "Mark selected items as shipped" in content
    # totals, payments, notes
    assert "Shipping</dt>" in content
    assert f"£{order.total:.2f}" in content
    assert f"pi_{order.pk}" in content
    assert "Gift wrap please" in content
    # actions: completing waits until everything has been shipped
    assert order_url(order, "complete") not in content
    assert '<button type="button" class="button button-secondary" disabled' in content
    assert (
        "The order can be marked completed once all its items have been marked "
        "shipped."
    ) in content
    assert order_url(order, "refund") in content
    assert order_url(order, "cancel") not in content
    assert resp.context["can_cancel"] is False


def test_order_page_all_shipped(admin_client, shop_variants, dispatcher):
    order = paid_order(shop_variants)
    order.mark_items_shipped(order.items.all(), dispatcher)
    resp = admin_client.get(order_url(order))
    content = resp.content.decode()
    assert order_url(order, "complete") in content
    assert "disabled" not in content.split("shop-order__actions")[1].split("</div>")[0]
    assert resp.context["can_ship"] is False


def test_order_page_collected(admin_client, shop_variants):
    order = paid_order(shop_variants, shipping_method="collect")
    content = admin_client.get(order_url(order)).content.decode()
    assert "Paid – ready to collect" in content
    assert "Ready to collect" in content
    assert "Mark selected items as collected" in content
    assert "Address" not in content


def test_order_page_unpaid(admin_client, shop_variants):
    order = make_order(shop_variants, status=OrderStatus.HOLD)
    resp = admin_client.get(order_url(order))
    content = resp.content.decode()
    assert "Awaiting payment" in content
    assert 'name="items"' not in content
    assert "Items can be marked shipped once the order has been paid." in content
    assert "Still to pay" in content
    assert "No payments yet." in content
    assert order_url(order, "cancel") in content
    assert order_url(order, "complete") not in content
    assert order_url(order, "refund") not in content


def test_order_page_discount_and_refunds(admin_client, shop_variants):
    mug, card, raffle = shop_variants
    basket = baker.make("shop.Basket", shipping_method="collect", extra={"email": "c@test.com"})
    basket.add(mug)
    basket.set_discount_code(baker.make("shop.DiscountCode", code="PINS5OFF", amount=5))
    order = Order.objects.create_from_basket(basket, None, status=OrderStatus.PROCESSING)
    order.pay(amount=order.total, transaction_id="pi_1", payment_method="stripe")
    order.pay(amount=Decimal("-2"), transaction_id="re_1", payment_method="stripe")
    content = admin_client.get(order_url(order)).content.decode()
    assert "PINS5OFF</dt>" in content
    assert "-£5.00" in content
    assert "£2.00 (partially refunded)" in content
    assert "Refund -£2.00" in content


def test_order_page_shipped_items(admin_client, shop_variants, dispatcher, freezer):
    freezer.move_to(datetime(2026, 10, 1, 12, tzinfo=UTC))
    order = paid_order(shop_variants)
    order.mark_items_shipped(order.items.all(), dispatcher)
    content = admin_client.get(order_url(order)).content.decode()
    assert "Shipped\n" in content and "1 Oct 2026" in content
    assert "by Ann Smith" in content
    assert 'name="items"' not in content


def test_order_page_not_found(admin_client):
    assert admin_client.get(reverse("shop_orders:edit", args=(999,))).status_code == 404


def test_order_add_note(admin_client, order):
    resp = admin_client.post(order_url(order), {"message": "Called the customer"})
    assert resp.status_code == 302
    assert [note.message for note in order.notes.all()] == ["Called the customer"]
    assert messages_for(resp) == ["Note added."]


def test_order_add_note_empty(admin_client, order):
    resp = admin_client.post(order_url(order), {"message": ""})
    assert resp.status_code == 200
    assert resp.context["note_form"].errors
    assert not order.notes.exists()


# Marking items shipped


def test_ship_selected_items(admin_client, shop_variants, django_capture_on_commit_callbacks):
    mug, card, raffle = shop_variants
    order = paid_order([mug, card])
    mug_item = order.items.get(product_id=mug.id)
    with django_capture_on_commit_callbacks(execute=True):
        resp = admin_client.post(order_url(order, "ship"), {"items": [mug_item.pk]})
    assert resp.url == order_url(order)
    assert messages_for(resp) == [
        "1 item marked shipped, and the customer has been emailed."
    ]
    mug_item.refresh_from_db()
    assert mug_item.shipping_status == OrderItem.SHIPPED
    assert mug_item.shipped_by.is_superuser
    order.refresh_from_db()
    assert order.status == OrderStatus.PARTIALLY_SHIPPED
    assert len(mail.outbox) == 1


def test_ship_all_items(admin_client, shop_variants):
    mug, card, raffle = shop_variants
    order = paid_order([mug, card])
    items = [item.pk for item in order.items_to_send]
    resp = admin_client.post(order_url(order, "ship"), {"items": items})
    assert not order.items_to_send.exists()
    order.refresh_from_db()
    assert order.status == OrderStatus.SHIPPED
    assert messages_for(resp) == [
        "2 items marked shipped, and the customer has been emailed."
    ]


def test_ship_collected_items(admin_client, shop_variants):
    order = paid_order(shop_variants, shipping_method="collect")
    items = [item.pk for item in order.items_to_send]
    resp = admin_client.post(order_url(order, "ship"), {"items": items})
    assert messages_for(resp) == [
        "2 items marked collected, and the customer has been emailed."
    ]


@pytest.mark.parametrize(
    "action,question,note,yes",
    [
        ("ship", "Mark these items as shipped?", "<strong>The customer will be emailed</strong> to say they've been shipped.", "Yes, mark as shipped"),
        ("not_sent", "Mark these items as won't be sent?", "The customer won't be emailed: contact them, and refund the items if necessary.", "Yes, mark as won't be sent"),
    ],
)
def test_ship_confirm(admin_client, shop_variants, action, question, note, yes):
    # first a page to confirm, listing the items: nothing's changed yet
    mug, card, raffle = shop_variants
    order = paid_order([mug, card])
    mug_item = order.items.get(product_id=mug.id)
    resp = admin_client.get(order_url(order, "ship"), {"items": [mug_item.pk], "action": action})
    assert resp.status_code == 200
    content = " ".join(html.unescape(resp.content.decode()).split())
    assert question in content
    assert "<ul> <li>2 x Test Product - Mug</li> </ul>" in content
    assert note in content
    assert f'<input type="hidden" name="items" value="{mug_item.pk}">' in content
    assert f'<input type="hidden" name="action" value="{action}">' in content
    assert yes in content
    assert f'<a href="{order_url(order)}" class="button button-secondary">No, go back</a>' in content
    mug_item.refresh_from_db()
    assert mug_item.shipping_status == OrderItem.TO_SEND
    assert not mail.outbox


def test_ship_confirm_collected(admin_client, shop_variants):
    order = paid_order(shop_variants, shipping_method="collect")
    items = [item.pk for item in order.items_to_send]
    content = html.unescape(admin_client.get(order_url(order, "ship"), {"items": items}).content.decode())
    assert "Mark these items as collected?" in content
    assert "Yes, mark as collected" in content


def test_ship_no_items_chosen(admin_client, shop_variants):
    order = paid_order(shop_variants)
    resp = admin_client.post(order_url(order, "ship"), {})
    assert messages_for(resp) == ["Choose the items first."]
    assert order.items_to_send.count() == 2


def test_ship_unpaid_order(admin_client, shop_variants):
    order = make_order(shop_variants, status=OrderStatus.HOLD)
    items = [item.pk for item in order.items_to_send]
    resp = admin_client.post(order_url(order, "ship"), {"items": items})
    assert messages_for(resp) == ["Items can't be marked shipped until the order is paid."]
    assert order.items_to_send.count() == 2


def test_ship_confirm_after_changed(admin_client, shop_variants, dispatcher):
    # e.g. back to the page to confirm after confirming: back to the order
    mug, card, raffle = shop_variants
    order = shipped_order([mug], dispatcher)
    item = order.items.get()
    resp = admin_client.get(order_url(order, "ship"), {"items": [item.pk], "action": "ship"})
    assert resp.url == order_url(order)
    assert messages_for(resp) == ["Those items have already been changed."]
    resp = admin_client.get(order_url(order, "ship"))
    assert messages_for(resp)[-1] == "Choose the items first."


# Completing and cancelling


def shipped_order(variants, user):
    order = paid_order(variants)
    order.mark_items_shipped(order.items.all(), user)
    mail.outbox.clear()
    return order


def test_complete_page(admin_client, shop_variants, dispatcher):
    order = shipped_order(shop_variants, dispatcher)
    content = html.unescape(admin_client.get(order_url(order, "complete")).content.decode())
    assert "Mark order <strong>" in content
    assert "The customer will be emailed" in content
    assert "Yes, mark as completed" in content


def test_complete_order(
    admin_client, shop_variants, dispatcher, django_capture_on_commit_callbacks
):
    order = shipped_order(shop_variants, dispatcher)
    with django_capture_on_commit_callbacks(execute=True):
        resp = admin_client.post(order_url(order, "complete"))
    order.refresh_from_db()
    assert order.status == OrderStatus.COMPLETED
    assert messages_for(resp) == [
        f"Order {order.ref} is completed, and the customer has been emailed."
    ]
    (email,) = mail.outbox
    assert email.subject == f"PINS shop order {order.ref} is complete"


def test_complete_not_allowed(admin_client, shop_variants):
    order = make_order(shop_variants, status=OrderStatus.HOLD)
    content = html.unescape(admin_client.get(order_url(order, "complete")).content.decode())
    assert "can't be changed: it's awaiting payment" in content
    resp = admin_client.post(order_url(order, "complete"))
    assert messages_for(resp) == [
        f"Order {order.ref} can't be changed: it's awaiting payment."
    ]
    order.refresh_from_db()
    assert order.status == OrderStatus.HOLD


@pytest.mark.parametrize("shipping_method,done", [("deliver", "shipped"), ("collect", "collected")])
def test_complete_before_shipping(admin_client, shop_variants, shipping_method, done):
    order = paid_order(shop_variants, shipping_method=shipping_method)
    message = (
        f"Order {order.ref} can't be marked completed until all its items have been "
        f"marked {done}."
    )
    content = html.unescape(admin_client.get(order_url(order, "complete")).content.decode())
    assert message in content
    assert "Yes, mark as completed" not in content
    resp = admin_client.post(order_url(order, "complete"))
    assert messages_for(resp) == [message]
    order.refresh_from_db()
    assert order.status == OrderStatus.PROCESSING


def test_cancel_order(admin_client, shop_variants):
    order = make_order(shop_variants, status=OrderStatus.HOLD)
    content = html.unescape(admin_client.get(order_url(order, "cancel")).content.decode())
    assert "The customer won't be emailed" in content
    resp = admin_client.post(order_url(order, "cancel"))
    order.refresh_from_db()
    assert order.status == OrderStatus.CANCELLED
    assert messages_for(resp) == [f"Order {order.ref} has been cancelled."]
    assert not mail.outbox


def test_action_breadcrumbs(admin_client, shop_variants):
    order = paid_order(shop_variants)
    content = admin_client.get(order_url(order, "complete")).content.decode()
    assert f'href="{INDEX_URL}"' in content
    assert f'href="{order_url(order)}"' in content
    assert "Mark as completed: order" in content


# Django admin (for tech admins)


def test_django_admin(admin_client, shop_variants):
    order = paid_order(shop_variants)
    resp = admin_client.get(reverse("admin:shop_order_changelist"))
    assert resp.status_code == 200
    resp = admin_client.get(reverse("admin:shop_order_change", args=(order.pk,)))
    assert resp.status_code == 200
    assert "Ann Smith" in resp.content.decode()


# Items that won't be sent


def test_status_items_not_sent(shop_variants, dispatcher):
    # some items shipped, the rest won't be sent: the order is shipped
    mug, card, raffle = shop_variants
    order = paid_order([mug, card, raffle])
    order.mark_items_shipped([order.items.get(product_id=mug.id)], dispatcher)
    changed = order.mark_items_not_sent(order.items.all(), dispatcher)
    assert [item.product_id for item in changed] == [card.id]
    card_item = order.items.get(product_id=card.id)
    assert card_item.shipping_status == OrderItem.NOT_SENT
    assert card_item.get_shipping_status_display() == "Won't be sent"
    assert card_item.shipped_by == dispatcher
    order.refresh_from_db()
    assert order.status == OrderStatus.SHIPPED
    assert order.can_complete
    # already changed
    assert order.mark_items_not_sent([card_item], dispatcher) == []


def test_status_no_items_sent(shop_variants, dispatcher):
    # nothing will be sent: the order stays processing, and can be completed
    order = paid_order(shop_variants)
    order.mark_items_not_sent(order.items.all(), dispatcher)
    order.refresh_from_db()
    assert order.status == OrderStatus.PROCESSING
    assert order.display_status == "Paid – nothing to send"
    assert order.can_complete


def test_shipped_email_ignores_items_not_sent(shop_variants, dispatcher, django_capture_on_commit_callbacks):
    mug, card, raffle = shop_variants
    order = paid_order([mug, card])
    order.mark_items_not_sent([order.items.get(product_id=card.id)], dispatcher)
    assert not mail.outbox
    with django_capture_on_commit_callbacks(execute=True):
        order.mark_items_shipped(order.items.all(), dispatcher)
    (email,) = mail.outbox
    assert email.subject == f"PINS shop order {order.ref} has been shipped"
    assert "Card" not in email.body


def test_mark_items_not_sent(admin_client, shop_variants):
    mug, card, raffle = shop_variants
    order = paid_order([mug, card])
    card_item = order.items.get(product_id=card.id)
    resp = admin_client.post(
        order_url(order, "ship"), {"items": [card_item.pk], "action": "not_sent"}
    )
    # the mug is still to send, so only the card needs refunding
    assert messages_for(resp) == [
        "1 item marked as won't be sent. The customer hasn't been emailed; contact "
        "them, and ask a Stripe admin to refund the items if necessary."
    ]
    card_item.refresh_from_db()
    assert card_item.shipping_status == OrderItem.NOT_SENT
    assert not mail.outbox
    content = html.unescape(admin_client.get(order_url(order)).content.decode())
    assert "Won't be sent" in content
    assert "Mark selected items as won't be sent" in content
    assert (
        "Mark items as won't be sent if they can't or won't be sent (e.g. out of "
        "stock or no longer wanted): the customer isn't emailed, so contact them, and "
        "refund the items if necessary. An entire order can be "
        "refunded from this page; to refund part of an order, ask a Stripe admin."
    ) in " ".join(content.split())
    assert content.count('name="items"') == 1


def test_mark_all_items_not_sent(admin_client, shop_variants):
    # nothing in the order is being sent: refund the order
    mug, card, raffle = shop_variants
    order = paid_order([mug, card])
    items = [item.pk for item in order.items.all()]
    resp = admin_client.post(order_url(order, "ship"), {"items": items, "action": "not_sent"})
    assert messages_for(resp) == [
        "2 items marked as won't be sent. The customer hasn't been emailed; contact "
        "them, and refund this order if necessary."
    ]


def test_mark_items_not_sent_with_unshipped_product(admin_client, shop_variants):
    # the raffle ticket doesn't need sending, so the order isn't refunded in full
    mug, card, raffle = shop_variants
    order = paid_order([mug, raffle])
    resp = admin_client.post(
        order_url(order, "ship"),
        {"items": [order.items.get(product_id=mug.id).pk], "action": "not_sent"},
    )
    assert "ask a Stripe admin to refund the items" in messages_for(resp)[0]


def test_customer_order_page_link(admin_client, order):
    # the order's status page, as the customer sees it (its address is encoded,
    # so admins can't easily find it otherwise)
    content = admin_client.get(reverse("shop_orders:edit", args=(order.pk,))).content.decode()
    url = reverse("shop:order_status", args=(order.token,))
    assert f'<a href="{url}" target="_blank" rel="noopener">Customer\'s order page' in content


# Undoing shipped (or collected) and won't be sent


def undo(admin_client, order, item, confirm=True):
    # confirm=False: the page to confirm
    method = admin_client.post if confirm else admin_client.get
    return method(order_url(order, "ship"), {"undo": item.pk})


def test_items_form(admin_client, shop_variants):
    # to the page to confirm (GET, so the back button works there)
    order = paid_order(shop_variants)
    content = admin_client.get(order_url(order)).content.decode()
    assert f'<form method="get" action="{order_url(order, "ship")}">' in content


def test_undo_buttons(admin_client, shop_variants, dispatcher):
    mug, card, raffle = shop_variants
    order = paid_order([mug, card])
    mug_item, card_item = order.items.get(product_id=mug.id), order.items.get(product_id=card.id)
    order.mark_items_shipped([mug_item], dispatcher)
    content = admin_client.get(order_url(order)).content.decode()
    assert f'name="undo" value="{mug_item.pk}"' in content
    assert f'name="undo" value="{card_item.pk}"' not in content
    # not once the order's completed
    order.mark_items_shipped([card_item], dispatcher)
    order.status = OrderStatus.COMPLETED
    order.save()
    assert 'name="undo"' not in admin_client.get(order_url(order)).content.decode()


def test_undo_shipped(admin_client, shop_variants, dispatcher, django_capture_on_commit_callbacks):
    mug, card, raffle = shop_variants
    order = shipped_order([mug, card], dispatcher)
    assert order.status == OrderStatus.SHIPPED
    mug_item = order.items.get(product_id=mug.id)
    # confirm first, saying the customer was emailed
    content = " ".join(html.unescape(undo(admin_client, order, mug_item, confirm=False).content.decode()).split())
    assert "Undo marking this item as shipped?" in content
    assert (
        "It will be ready to send again. The customer won't be emailed (they've "
        "already been emailed that it was shipped, so you may want to contact them)."
    ) in content
    assert "Yes, undo" in content
    assert f'<input type="hidden" name="undo" value="{mug_item.pk}">' in content
    mug_item.refresh_from_db()
    assert mug_item.shipping_status == OrderItem.SHIPPED
    # then: to send again, the order partly shipped, logged, and no emails (the
    # customer's or dispatchers' "paid" emails)
    with django_capture_on_commit_callbacks(execute=True):
        resp = undo(admin_client, order, mug_item)
    assert messages_for(resp) == ["Undone: 1 item is ready to send again. The customer hasn't been emailed."]
    mug_item.refresh_from_db()
    assert (mug_item.shipping_status, mug_item.shipped_at, mug_item.shipped_by) == (OrderItem.TO_SEND, None, None)
    order.refresh_from_db()
    assert order.status == OrderStatus.PARTIALLY_SHIPPED
    assert not mail.outbox
    content = html.unescape(admin_client.get(order_url(order)).content.decode())
    assert "Undone, ready to send again: 2 × Test Product - Mug" in content
    # undoing the other: back to being processed, still no emails
    with django_capture_on_commit_callbacks(execute=True):
        undo(admin_client, order, order.items.get(product_id=card.id))
    order.refresh_from_db()
    assert order.status == OrderStatus.PROCESSING
    assert order.display_status == "Paid – ready to ship"
    assert not mail.outbox


def test_undo_not_sent(admin_client, shop_variants, dispatcher):
    # the card won't be sent and the mug's shipped (so the order's shipped);
    # undoing the card: partly shipped
    mug, card, raffle = shop_variants
    order = paid_order([mug, card], shipping_method="collect")
    card_item = order.items.get(product_id=card.id)
    order.mark_items_not_sent([card_item], dispatcher)
    order.mark_items_shipped([order.items.get(product_id=mug.id)], dispatcher)
    order.refresh_from_db()
    assert order.status == OrderStatus.SHIPPED
    content = " ".join(html.unescape(undo(admin_client, order, card_item, confirm=False).content.decode()).split())
    assert "Undo marking this item as won't be sent?" in content
    assert "It will be ready to collect again. The customer won't be emailed." in content
    resp = undo(admin_client, order, card_item)
    assert messages_for(resp) == ["Undone: 1 item is ready to collect again. The customer hasn't been emailed."]
    order.refresh_from_db()
    assert order.status == OrderStatus.PARTIALLY_SHIPPED
    assert "Undone, ready to collect again: 2 × Card" in html.unescape(admin_client.get(order_url(order)).content.decode())


def test_undo_not_allowed(admin_client, shop_variants, dispatcher):
    # once the order's completed (or refunded, cancelled), or for an item still to send
    mug, card, raffle = shop_variants
    order = shipped_order([mug], dispatcher)
    item = order.items.get()
    order.status = OrderStatus.COMPLETED
    order.save()
    resp = undo(admin_client, order, item)
    assert messages_for(resp) == ["Items can't be changed: the order is completed."]
    assert item.shipping_status == OrderItem.SHIPPED
    assert order.mark_items_to_send([item], dispatcher) == []
    order = paid_order([mug])
    resp = undo(admin_client, order, order.items.get())
    assert messages_for(resp)[-1] == "Those items have already been changed."
