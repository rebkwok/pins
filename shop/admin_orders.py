"""
The orders admin: a list of orders (with filters and a spreadsheet export of
their items) and an order page with actions (mark items shipped, complete,
cancel, refund, add a note), for shop admins. This replaces salesman's Wagtail
orders admin; tech admins can check orders' fields in the Django admin.
"""

import django_filters
from django import forms
from django.contrib.auth import get_user_model
from django.contrib.contenttypes.models import ContentType
from django.core.exceptions import PermissionDenied
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect
from django.template.defaultfilters import pluralize
from django.urls import path, reverse
from django.utils.dateformat import format as date_format
from django.utils.functional import cached_property
from django.utils.timezone import localtime
from django.utils.html import format_html
from django.views.generic import TemplateView
from salesman.core.utils import get_salesman_model
from wagtail.admin import messages
from wagtail.admin.filters import DateRangePickerWidget, WagtailFilterSet
from wagtail.admin.ui.components import Component
from wagtail.admin.ui.tables import BaseColumn, Column, TitleColumn
from wagtail.admin.views.generic import IndexView
from wagtail.admin.views.generic.base import WagtailAdminTemplateMixin
from wagtail.admin.views.mixins import SpreadsheetExportMixin
from wagtail.admin.viewsets.model import ModelViewSet
from wagtail.log_actions import registry as log_registry
from wagtail.permission_policies import ModelPermissionPolicy

from common.users import real_users_q

from . import activity
from .admin_views import SHIPPABLE_STATUSES
from .models import DISPATCHERS_GROUP, Order, OrderItem, Product, ProductVariant
from .payment import PayByStripe
from .templatetags.shoptags import gbp

User = get_user_model()
OrderNote = get_salesman_model("OrderNote")

URL_NAMESPACE = "shop_orders"

# Status filter: the statuses actually used, in plain words
AWAITING_PAYMENT = [Order.Status.NEW, Order.Status.CREATED, Order.Status.HOLD]
STATUS_FILTER_CHOICES = {
    "awaiting_payment": ("Awaiting payment", AWAITING_PAYMENT),
    "failed": ("Payment failed", [Order.Status.FAILED]),
    "processing": ("Paid – ready to ship/collect", [Order.Status.PROCESSING]),
    "partly_shipped": ("Partly shipped/collected", [Order.Status.PARTIALLY_SHIPPED]),
    "shipped": ("Shipped/collected", [Order.Status.SHIPPED]),
    "completed": ("Completed", [Order.Status.COMPLETED]),
    "cancelled": ("Cancelled", [Order.Status.CANCELLED]),
    "refunded": ("Refunded", [Order.Status.REFUNDED]),
}
REFUND_FILTER_CHOICES = {
    "partial": "Partially refunded",
    "full": "Fully refunded",
    "none": "",
}


def _variant_ids(product):
    return ProductVariant.objects.filter(product=product).values("id")


def _variant_items(product):
    """A Q for order items of the product's variants"""
    return Q(
        product_content_type=ContentType.objects.get_for_model(ProductVariant),
        product_id__in=_variant_ids(product),
    )


class OrderFilterSet(WagtailFilterSet):
    status = django_filters.ChoiceFilter(
        label="Status",
        choices=[(key, label) for key, (label, _) in STATUS_FILTER_CHOICES.items()],
        method="filter_status",
        empty_label="Any status",
    )
    to_ship = django_filters.BooleanFilter(
        label="Has items to ship or collect",
        method="filter_to_ship",
        widget=forms.Select(
            choices=[("", "Any"), ("true", "Yes"), ("false", "No")],
        ),
    )
    dispatcher = django_filters.ModelChoiceFilter(
        label="Dispatcher",
        queryset=User.objects.none(),
        method="filter_dispatcher",
        empty_label="Any dispatcher",
    )
    # (from Dispatcher activity: the items that shop admins send themselves)
    no_dispatcher = django_filters.BooleanFilter(
        label="Has items without a dispatcher to ship or collect",
        method="filter_no_dispatcher",
        widget=forms.Select(choices=[("", "Any"), ("true", "Yes")]),
    )
    product = django_filters.ModelChoiceFilter(
        label="Product",
        queryset=Product.objects.order_by("category__sort_order", "category__title", "name"),
        method="filter_product",
        empty_label="Any product",
    )
    refunds = django_filters.ChoiceFilter(
        label="Refunds",
        choices=[(key, label or "Not refunded") for key, label in REFUND_FILTER_CHOICES.items()],
        method="filter_refunds",
        empty_label="Any",
    )
    date_created = django_filters.DateFromToRangeFilter(
        label="Date", widget=DateRangePickerWidget
    )

    class Meta:
        model = Order
        fields = []

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        # Current dispatchers, and anyone who has been a dispatcher for an order
        self.filters["dispatcher"].queryset = (
            User.objects.filter(
                Q(groups__name=DISPATCHERS_GROUP) | Q(dispatched_order_items__isnull=False)
            )
            .filter(real_users_q())
            .distinct()
            .order_by("first_name", "last_name", "username")
        )
        self.filters["dispatcher"].field.label_from_instance = (
            lambda user: user.get_full_name() or user.get_username()
        )

    def filter_status(self, queryset, name, value):
        return queryset.filter(status__in=STATUS_FILTER_CHOICES[value][1])

    def filter_to_ship(self, queryset, name, value):
        to_ship = Q(
            items__shipping_status=OrderItem.TO_SEND, status__in=SHIPPABLE_STATUSES
        )
        ids = Order.objects.filter(to_ship).values("id")
        return queryset.filter(id__in=ids) if value else queryset.exclude(id__in=ids)

    def filter_no_dispatcher(self, queryset, name, value):
        if not value:
            return queryset
        return queryset.filter(
            id__in=Order.objects.filter(
                items__shipping_status=OrderItem.TO_SEND,
                items__dispatcher__isnull=True,
                status__in=SHIPPABLE_STATUSES,
            ).values("id")
        )

    def filter_dispatcher(self, queryset, name, value):
        return queryset.filter(
            id__in=OrderItem.objects.filter(dispatcher=value).values("order_id")
        )

    def filter_product(self, queryset, name, value):
        return queryset.filter(
            id__in=OrderItem.objects.filter(_variant_items(value)).values("order_id")
        )

    def filter_refunds(self, queryset, name, value):
        status = REFUND_FILTER_CHOICES[value]
        return queryset.filter(
            id__in=[
                order.id
                for order in queryset.prefetch_related("payments")
                if order.refund_status == status
            ]
        )


class OrderDownloadButton(Component):
    """Download the ticked orders (see order_export.js)"""

    template_name = "shop/admin/orders/download_button.html"


class SelectColumn(BaseColumn):
    """Checkboxes to choose the orders to download (see order_export.js)"""

    header_template_name = "shop/admin/orders/tables/select_header.html"
    cell_template_name = "shop/admin/orders/tables/select_cell.html"

    class Media:
        js = ["shop/js/order_export.js"]


class ItemsColumn(BaseColumn):
    """The order's items, collapsed ("3 items"), with their shipping status"""

    cell_template_name = "shop/admin/orders/tables/items_cell.html"


def customer_display(order):
    return format_html(
        "{}<br><small>{}</small>", order.name or "-", order.email
    )


class OrderItemsExportMixin(SpreadsheetExportMixin):
    """
    A spreadsheet of order items (one row per item), with each order's rows
    shaded in turn in Excel (e.g. the orders list's download, and dispatchers')
    """

    # Each order's rows (one per item) are shaded in turn, so the orders can be
    # told apart: the first order's rows grey, the next's white, and so on. All
    # the cells have borders, and the headings are bold. Columns are as wide as
    # their longest line, up to MAX_COLUMN_WIDTH (characters); longer values,
    # and values on several lines (e.g. addresses), wrap.
    ORDER_SHADING = "D9D9D9"
    BORDER_COLOUR = "A6A6A6"
    MAX_COLUMN_WIDTH = 50
    # (dates, as Excel shows them)
    DATE_WIDTH = 16
    export_sheet_title = "Orders"

    @classmethod
    def cell_width(cls, value):
        """The width a cell's value needs: its longest line"""
        if value is None:
            return 0
        if hasattr(value, "year"):
            return cls.DATE_WIDTH
        return max(len(line) for line in str(value).split("\n"))

    def write_xlsx(self, queryset, output):
        # (queryset: order items, in their orders' order)
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from wagtail.admin.views.mixins import ExcelDateFormatter

        workbook = Workbook(write_only=True, iso_dates=True)
        worksheet = workbook.create_sheet(title=self.export_sheet_title)
        side = Side(style="thin", color=self.BORDER_COLOUR)
        border = Border(left=side, right=side, top=side, bottom=side)
        headings = []
        for field in self.list_export:
            cell = WriteOnlyCell(worksheet, self.get_heading(queryset, field))
            cell.font = Font(bold=True)
            cell.border = border
            headings.append(cell)
        rows = [headings]
        date_format = ExcelDateFormatter().get()
        shading = PatternFill("solid", fgColor=self.ORDER_SHADING)
        order_id, shaded = None, False
        for item in queryset:
            if item.order_id != order_id:
                order_id, shaded = item.order_id, not shaded
            cells = list(
                self.generate_xlsx_row(
                    worksheet, self.to_row_dict(item), date_format=date_format
                )
            )
            for cell in cells:
                cell.border = border
                if shaded:
                    cell.fill = shading
            rows.append(cells)
        # (set before the rows are written, in a write-only worksheet)
        wrap = Alignment(wrap_text=True, vertical="top")
        for column, cells in enumerate(zip(*rows), start=1):
            width = max(self.cell_width(cell.value) for cell in cells)
            worksheet.column_dimensions[get_column_letter(column)].width = (
                min(width, self.MAX_COLUMN_WIDTH) + 2
            )
            for cell in cells:
                value = cell.value
                if isinstance(value, str) and (
                    "\n" in value or len(value) > self.MAX_COLUMN_WIDTH
                ):
                    cell.alignment = wrap
        for cells in rows:
            worksheet.append(cells)
        workbook.save(output)


class OrderIndexView(OrderItemsExportMixin, IndexView):
    # (e.g. a search from the sidebar can be run in the other areas, e.g. products)
    show_other_searches = True

    # Dates in CSV files as e.g. "2026-10-24 14:05" (Excel files have dates)
    custom_field_preprocess = {
        field: {
            IndexView.FORMAT_CSV: lambda value: (
                f"{localtime(value):%Y-%m-%d %H:%M}" if value else ""
            )
        }
        for field in ["order.date_created", "shipped_at"]
    }

    def search_queryset(self, queryset):
        # The order's reference, customer and email, or its items' names (e.g.
        # orders that included a hoodie)
        if not self.search_query:
            return queryset
        # (items' names are in their product data)
        items = OrderItem.objects.filter(product_data__name__icontains=self.search_query)
        return super().search_queryset(queryset) | queryset.filter(
            id__in=items.values("order_id")
        )

    def get_base_queryset(self):
        # display_status uses the payments; the items column, the items
        return super().get_base_queryset().prefetch_related("payments", "items")

    # Export: one row per order item, of the ticked orders, or all the (filtered)
    # orders

    # The download button is in the header, beside the title (rather than in
    # the header's "..." menu), with help above the results
    results_template_name = "shop/admin/orders/index_results.html"
    header_more_buttons = []

    @cached_property
    def header_buttons(self):
        # (orders can't be added)
        return [OrderDownloadButton()]

    def get_filename(self):
        return f"orders-{localtime():%Y-%m-%d-%H%M}"

    def as_spreadsheet(self, queryset, spreadsheet_format):
        # The ticked orders, or (with "select all") all the orders, with any
        # filters
        ticked = self.request.GET.getlist("order")
        if ticked:
            queryset = queryset.filter(pk__in=[pk for pk in ticked if pk.isdigit()])
        items = (
            OrderItem.objects.filter(order__in=queryset)
            .select_related("order")
            .order_by("-order__date_created", "order_id", "id")
        )
        # Only the chosen dispatcher's or product's items
        filters = self.filters.form.cleaned_data if self.filters else {}
        if filters.get("dispatcher"):
            items = items.filter(dispatcher=filters["dispatcher"])
        if filters.get("product"):
            items = items.filter(_variant_items(filters["product"]))
        return super().as_spreadsheet(items, spreadsheet_format)


    def get_heading(self, queryset, field):
        return self.export_headings[field]


class OrderPermissionPolicy(ModelPermissionPolicy):
    """Orders are created by checkout, and never deleted"""

    def user_has_permission(self, user, action):
        if action in ("add", "delete"):
            return False
        return super().user_has_permission(user, action)

    def user_has_any_permission(self, user, actions):
        return any(self.user_has_permission(user, action) for action in actions)


def stripe_payments(order):
    return order.payments.filter(payment_method=PayByStripe.identifier, amount__gt=0)


class OrderViewMixin(WagtailAdminTemplateMixin):
    """An order's admin page (or action), for users who can change orders"""

    header_icon = "list-ul"

    def dispatch(self, request, *args, **kwargs):
        if not OrderPermissionPolicy(Order).user_has_permission(request.user, "change"):
            raise PermissionDenied
        return super().dispatch(request, *args, **kwargs)

    @cached_property
    def order(self):
        return get_object_or_404(Order, pk=self.kwargs["pk"])

    @property
    def order_url(self):
        return reverse(f"{URL_NAMESPACE}:edit", args=(self.order.pk,))

    def get_page_title(self):
        return f"Order {self.order.ref}"

    def get_breadcrumbs_items(self):
        return self.breadcrumbs_items + [
            {"url": reverse(f"{URL_NAMESPACE}:index"), "label": "Orders"},
            {"url": self.order_url, "label": f"Order {self.order.ref}"},
        ]

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["order"] = self.order
        context["order_url"] = self.order_url
        return context


class OrderActionMixin(OrderViewMixin):
    """A page for one of the order's actions (e.g. confirming a refund)"""

    action_label = None

    def get_page_title(self):
        return f"{self.action_label}: order {self.order.ref}"

    def get_breadcrumbs_items(self):
        return super().get_breadcrumbs_items() + [{"url": "", "label": self.action_label}]

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["action_label"] = self.action_label
        return context


class OrderNoteForm(forms.Form):
    message = forms.CharField(
        label="Add a note",
        widget=forms.Textarea(attrs={"rows": 3}),
        help_text="Notes are only seen by admins, not the customer.",
    )


class OrderDetailView(OrderViewMixin, TemplateView):
    """The order page"""

    template_name = "shop/admin/orders/detail.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        order = self.order
        items = list(order.items.select_related("dispatcher", "shipped_by"))
        context.update(
            items=items,
            can_ship=order.can_ship_items
            and any(item.shipping_status == OrderItem.TO_SEND for item in items),
            can_undo=order.can_undo_items,
            # (its items to send won't be)
            closed=order.status
            in (Order.Status.REFUNDED, Order.Status.CANCELLED, Order.Status.FAILED),
            payments=order.payments.order_by("date_created"),
            notes=order.notes.order_by("date_created"),
            # who did what to the order (activity.py), oldest first
            activity=log_registry.get_logs_for_instance(order)
            .select_related("user")
            .order_by("timestamp", "pk"),
            note_form=kwargs.get("note_form") or OrderNoteForm(),
            can_complete=order.can_complete,
            complete_after_shipping=order.can_change_status(Order.Status.COMPLETED)
            and not order.can_complete,
            can_cancel=order.can_change_status(Order.Status.CANCELLED),
            can_refund=order.status != Order.Status.REFUNDED
            and stripe_payments(order).exists(),
        )
        return context

    def post(self, request, *args, **kwargs):
        # Add a note
        form = OrderNoteForm(request.POST)
        if not form.is_valid():
            return self.render_to_response(self.get_context_data(note_form=form))
        # (not order.notes.create(): notes are a modelcluster relation, which
        # doesn't save)
        message = form.cleaned_data["message"]
        OrderNote.objects.create(order=self.order, message=message)
        activity.log_order(self.order, activity.ORDER_NOTE, message=message)
        messages.success(request, "Note added.")
        return redirect(self.order_url)


class OrderShipView(OrderActionMixin, TemplateView):
    """
    Mark items shipped (or collected), won't be sent, or ready to send again
    (undo), from the order page: its form opens this page (GET, so the browser's
    back button can come back to it) to confirm, as marking items shipped
    emails the customer; confirming posts the same fields, then it's back to
    the order page.
    """

    template_name = "shop/admin/confirm_items.html"
    action_label = "Confirm"

    @property
    def data(self):
        return self.request.POST if self.request.method == "POST" else self.request.GET

    @property
    def action(self):
        if "undo" in self.data:
            return "undo"
        return "not_sent" if self.data.get("action") == "not_sent" else "ship"

    def get_items(self):
        order = self.order
        if self.action == "undo":
            return list(
                order.items.filter(
                    pk=self.data.get("undo"),
                    shipping_status__in=[OrderItem.SHIPPED, OrderItem.NOT_SENT],
                )
            )
        return list(order.items_to_send.filter(pk__in=self.data.getlist("items")))

    def confirm_context(self, items):
        order = self.order
        done = "collected" if order.is_collected else "shipped"
        fields = [("items", item.pk) for item in items]
        if self.action == "ship":
            question = f"Mark these items as {done}?"
            notes = [
                format_html(
                    "<strong>The customer will be emailed</strong> to say they've been {}.",
                    done,
                )
            ]
            fields.append(("action", "ship"))
            yes_label = f"Yes, mark as {done}"
        elif self.action == "not_sent":
            question = "Mark these items as won't be sent?"
            notes = [
                "The customer won't be emailed: contact them, and refund the items if "
                "necessary."
            ]
            fields.append(("action", "not_sent"))
            yes_label = "Yes, mark as won't be sent"
        else:
            ready = "ready to collect" if order.is_collected else "ready to send"
            item = items[0]
            was = done if item.shipping_status == OrderItem.SHIPPED else "won't be sent"
            question = f"Undo marking this item as {was}?"
            notes = [
                f"It will be {ready} again. The customer won't be emailed"
                + (
                    f" (they've already been emailed that it was {done}, so you may "
                    "want to contact them)."
                    if item.shipping_status == OrderItem.SHIPPED
                    else "."
                )
            ]
            fields = [("undo", item.pk)]
            yes_label = "Yes, undo"
        return {
            "question": question,
            "items": items,
            "notes": notes,
            "fields": fields,
            "yes_label": yes_label,
            "back_url": self.order_url,
        }

    def check(self):
        """The items to change, or a redirect back to the order page saying why not"""
        order = self.order
        allowed = order.can_undo_items if self.action == "undo" else order.can_ship_items
        if not allowed:
            if self.action == "undo":
                messages.error(
                    self.request,
                    f"Items can't be changed: the order is {order.display_status.lower()}.",
                )
            else:
                messages.error(
                    self.request, "Items can't be marked shipped until the order is paid."
                )
            return None, redirect(self.order_url)
        items = self.get_items()
        if not items:
            if "undo" in self.data or self.data.getlist("items"):
                # e.g. back to this page after confirming
                messages.warning(self.request, "Those items have already been changed.")
            else:
                messages.warning(self.request, "Choose the items first.")
            return None, redirect(self.order_url)
        return items, None

    def get(self, request, *args, **kwargs):
        items, response = self.check()
        if response:
            return response
        return self.render_to_response(self.get_context_data(**self.confirm_context(items)))

    def post(self, request, *args, **kwargs):
        order = self.order
        items, response = self.check()
        if response:
            return response
        count = f"{len(items)} item{pluralize(len(items))}"
        if self.action == "undo":
            changed = order.mark_items_to_send(items, request.user)
            ready = "ready to collect" if order.is_collected else "ready to send"
            messages.success(
                request,
                f"Undone: {len(changed)} item{pluralize(len(changed), ' is,s are')} {ready} again. "
                "The customer hasn't been emailed.",
            )
        elif self.action == "not_sent":
            order.mark_items_not_sent(items, request.user)
            if order.items.exclude(shipping_status=OrderItem.NOT_SENT).exists():
                # Some of the order is being sent (or doesn't need sending), so
                # only these items' payment needs refunding (in Stripe)
                refund = "and ask a Stripe admin to refund the items if necessary."
            else:
                refund = "and refund this order if necessary."
            messages.success(
                request,
                f"{count} marked as won't be sent. The customer hasn't been emailed; "
                f"contact them, {refund}",
            )
        else:
            order.mark_items_shipped(items, request.user)
            done = "collected" if order.is_collected else "shipped"
            messages.success(
                request, f"{count} marked {done}, and the customer has been emailed."
            )
        return redirect(self.order_url)


class OrderStatusView(OrderActionMixin, TemplateView):
    """Confirm, then change the order's status (e.g. complete or cancel it)"""

    template_name = "shop/admin/orders/confirm_status.html"
    new_status = None
    done_message = None

    def allowed(self):
        return self.order.can_change_status(self.new_status)

    def not_allowed_message(self):
        order = self.order
        return f"Order {order.ref} can't be changed: it's {order.display_status.lower()}."

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["allowed"] = self.allowed()
        context["not_allowed_message"] = self.not_allowed_message()
        return context

    def post(self, request, *args, **kwargs):
        order = self.order
        if not self.allowed():
            messages.error(request, self.not_allowed_message())
            return redirect(self.order_url)
        order.status = self.new_status
        order.save()
        activity.log_order(order, self.log_action)
        messages.success(request, self.done_message.format(order=order))
        return redirect(self.order_url)


class OrderCompleteView(OrderStatusView):
    new_status = Order.Status.COMPLETED
    log_action = activity.ORDER_COMPLETED
    action_label = "Mark as completed"
    done_message = "Order {order.ref} is completed, and the customer has been emailed."

    def allowed(self):
        # Once everything has been shipped or collected
        return self.order.can_complete

    def not_allowed_message(self):
        order = self.order
        if order.can_change_status(self.new_status):
            done = "collected" if order.is_collected else "shipped"
            return (
                f"Order {order.ref} can't be marked completed until all its items "
                f"have been marked {done}."
            )
        return super().not_allowed_message()


class OrderCancelView(OrderStatusView):
    new_status = Order.Status.CANCELLED
    log_action = activity.ORDER_CANCELLED
    action_label = "Cancel order"
    done_message = "Order {order.ref} has been cancelled."


class OrderRefundView(OrderActionMixin, TemplateView):
    """
    Refund an order's Stripe payments in full (partial refunds are made in the
    Stripe dashboard, and recorded by the webhook)
    """

    template_name = "shop/admin/orders/refund.html"
    action_label = "Refund"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["stripe_payments"] = stripe_payments(self.order)
        context["refund_amount"] = self.order.amount_paid
        # (which can be marked won't be sent as well)
        context["items_to_send"] = list(self.order.items_to_send)
        return context

    def post(self, request, *args, **kwargs):
        order = self.order
        if order.status == order.Status.REFUNDED:
            messages.error(request, f"Order {order} has already been refunded.")
            return redirect(self.order_url)
        # (before refunding, which may change the order's status)
        items_to_send = list(order.items_to_send)
        refunded, errors = PayByStripe.refund_order(order)
        if refunded:
            not_sent = ""
            if items_to_send and request.POST.get("mark_not_sent"):
                order.mark_items_not_sent(items_to_send, request.user)
                count = len(items_to_send)
                not_sent = f" {count} item{pluralize(count)} marked as won't be sent."
            messages.success(
                request,
                f"{gbp(refunded)} refunded for order {order}.{not_sent} Stock hasn't "
                "been changed; update it manually if items are going back into stock.",
            )
        for error in errors:
            messages.error(request, f"Error refunding order {order}: {error}")
        if not refunded and not errors:
            messages.warning(request, f"Order {order} has no Stripe payments to refund.")
        return redirect(self.order_url)


class OrderViewSet(ModelViewSet):
    model = Order
    name = URL_NAMESPACE
    url_prefix = "shop/orders"
    icon = "list-ul"
    menu_label = "Orders"
    # (the shop's menus are together, see dashboard/wagtail_hooks.py)
    menu_order = 510
    add_to_admin_menu = True

    index_view_class = OrderIndexView
    list_display = [
        SelectColumn("select"),
        # (Wagtail only links the first column to the order page)
        TitleColumn(
            "ref",
            label="Reference",
            sort_key="ref",
            get_url=lambda order: reverse(f"{URL_NAMESPACE}:edit", args=(order.pk,)),
        ),
        Column(
            "date_created",
            label="Date",
            accessor=lambda order: date_format(
                localtime(order.date_created), "j M Y, H:i"
            ),
        ),
        Column("customer", label="Customer", accessor=customer_display),
        ItemsColumn("items", label="Items"),
        Column("total", label="Total", accessor=lambda order: gbp(order.total)),
        Column("status", label="Status", accessor="display_status"),
    ]
    filterset_class = OrderFilterSet
    # (and the orders' items, e.g. "hoodie": see OrderIndexView.search_queryset)
    search_fields = ["ref", "name", "email"]
    ordering = ["-date_created"]
    list_per_page = 50

    list_export = [
        "order.ref",
        "order.date_created",
        "order.name",
        "order.email",
        "order.phone",
        "order.get_shipping_method_display",
        "order.shipping_address",
        "order.safe_place",
        "order.display_status",
        "name",
        "quantity",
        "total",
        "dispatcher_name",
        "status_label",
        "shipped_at",
    ]
    export_headings = {
        "order.ref": "Order",
        "order.date_created": "Date",
        "order.name": "Customer",
        "order.email": "Email",
        "order.get_shipping_method_display": "Shipping method",
        "order.phone": "Mobile number",
        "order.shipping_address": "Shipping address",
        "order.safe_place": "Safe place",
        "order.display_status": "Order status",
        "name": "Item",
        "quantity": "Quantity",
        "total": "Item total (£)",
        "dispatcher_name": "Dispatcher",
        "status_label": "Item status",
        "shipped_at": "Shipped/collected",
    }
    export_filename = "orders"

    @cached_property
    def permission_policy(self):
        return OrderPermissionPolicy(self.model)

    def get_urlpatterns(self):
        # No add, edit form, delete, history etc: the "edit" page is the order
        # page, with its actions
        return [
            path("", self.index_view, name="index"),
            path("results/", self.index_results_view, name="index_results"),
            path("<int:pk>/", OrderDetailView.as_view(), name="edit"),
            path("<int:pk>/ship/", OrderShipView.as_view(), name="ship"),
            path("<int:pk>/complete/", OrderCompleteView.as_view(), name="complete"),
            path("<int:pk>/cancel/", OrderCancelView.as_view(), name="cancel"),
            path("<int:pk>/refund/", OrderRefundView.as_view(), name="refund"),
        ]

    def register_reference_index(self):
        # Orders aren't referenced by anything shown in the admin
        pass
